"""과제 '활용 데이터' 엑셀 파일을 만든다.

과제 안내: "AI로 생성한 '직급별 업무 시간 할당량' 또는 '가장 하기 싫은 업무 설문조사' 더미 데이터"
- 두 데이터를 모두 담는다. 원자료는 P1이 쓰는 CSV와 같은 값이다(seed/dashboard/).
- 합계·비중·순위는 엑셀 수식으로 계산한다(값을 바꾸면 다시 계산된다).
- 입력값(더미)은 파란 글자, 수식은 검은 글자로 구분한다.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"
BLUE = "0000FF"
HEAD_FILL = PatternFill("solid", fgColor="E8EEF7")
SUB_FILL = PatternFill("solid", fgColor="F3F3F0")
THIN = Side(style="thin", color="C3C2B7")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STAGES = ["처리 파악", "증빙 준비", "문서 작성", "반려 후 재작성", "진행 확인·문의", "결재 검토"]
OTHERS = ["본업", "회의", "기타 사무"]
GRADES = ["사원", "선임", "책임", "수석", "팀장 이상"]
OPTIONS = STAGES + ["기타 사무"]
STAGE_DESC = {
    "처리 파악": "어떤 결재를 올려야 하는지, 양식·전결권자가 무엇인지 찾는 시간",
    "증빙 준비": "영수증 정리·옮겨 적기, 합계 대조, 중복·한도 확인",
    "문서 작성": "양식 항목 채우기, 메모 내용 옮기기",
    "반려 후 재작성": "누락·한도 초과·증빙 문제로 반려되어 다시 쓰는 시간",
    "진행 확인·문의": "결재가 어디까지 갔는지 확인하고 결재자에게 묻는 시간",
    "결재 검토": "결재자가 대기 결재를 고르고 문서마다 규정·증빙을 확인하는 시간",
}


def _f(bold=False, color=None, size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def _head(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font, c.fill, c.border = _f(bold=True), HEAD_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths: dict[str, float]):
    for k, w in widths.items():
        ws.column_dimensions[k].width = w


def build_workbook(path: Path, alloc: list[dict], survey: list[dict], seed: int, demo_start: str) -> None:
    wb = Workbook()
    _readme(wb.active, seed, demo_start, len(survey))
    _allocation(wb.create_sheet("직급별_업무시간_할당량"), alloc)
    _questions(wb.create_sheet("설문_문항"))
    _responses(wb.create_sheet("설문_응답"), survey)
    _summary(wb.create_sheet("설문_집계"), len(survey))
    wb.calculation.fullCalcOnLoad = True
    fixed = dt.datetime.fromisoformat(demo_start)
    wb.properties.creator = "approval-demo seedgen"
    wb.properties.created = fixed
    wb.properties.modified = fixed
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _readme(ws, seed: int, demo_start: str, n: int) -> None:
    ws.title = "설명"
    rows = [
        ("과제 활용 데이터 (더미)", None),
        ("", None),
        ("과제 조건", "AI로 생성한 '직급별 업무 시간 할당량' 또는 '가장 하기 싫은 업무 설문조사' 더미 데이터"),
        ("이 파일", "두 데이터를 모두 담았다. 모든 값은 시연용 더미이며 실제 회사 자료가 아니다."),
        ("생성 방법", "AI(Claude)가 작성한 생성기(seedgen)가 정한 가정과 고정 시드로 만든다. 같은 시드면 같은 값이 나온다."),
        ("시드", seed),
        ("기준 시점(시연 시계)", demo_start),
        ("다시 만들기", "uv run approval-seed build  →  seed/과제_활용데이터.xlsx"),
        ("", None),
        ("시트", "내용"),
        ("직급별_업무시간_할당량", "직급군 × 업무 범주(전결 처리는 6단계로 나눔)의 주당 시간. 합계·비중은 수식"),
        ("설문_문항", "설문 문항과 선택지 정의"),
        ("설문_응답", f"응답 {n}건 원자료 (부서, 직급군, 1~3순위, 주당 전결 처리 시간, 이유)"),
        ("설문_집계", "선택지별 1·2·3순위 수와 가중 점수(1순위 3점, 2순위 2점, 3순위 1점), 순위, 직급군별 1순위 교차표 (모두 수식)"),
        ("", None),
        ("업무 범주", "본업 · 회의 · 전결 처리 · 기타 사무 (설계서 8.2)"),
        ("전결 처리 6단계", " · ".join(STAGES) + " (설계서 1.2)"),
    ]
    for st in STAGES:
        rows.append((f"  {st}", STAGE_DESC[st]))
    rows += [
        ("", None),
        ("가정 (생성기에 적힌 값)", None),
        ("할당량", "직급군별 기본값에 ±5% 잡음. 주당 약 40시간. 결재 검토 시간은 팀장 이상에서 크게 두었다."),
        ("설문", "직급군별 선택 가중치를 다르게 둔 무작위 응답. 결재 검토는 팀장 이상에서 자주 고르도록 했다."),
        ("주의", "특정 결론(예: 하루 1시간 절감)에 맞추려고 값을 고치지 않았다. 실제 비중은 파일럿에서 측정한다."),
        ("표기", "파란 글자 = 생성된 입력값(더미), 검은 글자 = 수식"),
    ]
    for r, (k, v) in enumerate(rows, 1):
        a = ws.cell(row=r, column=1, value=k)
        a.font = _f(bold=r == 1 or k in ("시트", "가정 (생성기에 적힌 값)"), size=14 if r == 1 else 10)
        if v is not None:
            b = ws.cell(row=r, column=2, value=v)
            b.font = _f(bold=k == "시트")
            b.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, {"A": 26, "B": 96})


def _allocation(ws, alloc: list[dict]) -> None:
    val = {(r["category"], r["stage"], r["grade_group"]): float(r["hours_per_week"]) for r in alloc}
    ws["A1"] = "직급별 업무 시간 할당량 (주당 시간, 더미)"
    ws["A1"].font = _f(bold=True, size=13)
    ws["A2"] = "파란 글자는 생성된 입력값, 검은 글자는 수식. 단위: 시간/주"
    ws["A2"].font = _f(color="52514E")
    _head(ws, 4, ["업무 범주", "세부 단계"] + GRADES)
    r = 5
    first = r
    stage_rows = []
    for cat in OTHERS:
        ws.cell(row=r, column=1, value=cat).font = _f()
        ws.cell(row=r, column=2, value="—").font = _f(color="898781")
        for j, g in enumerate(GRADES):
            c = ws.cell(row=r, column=3 + j, value=val.get((cat, "", g), 0.0))
            c.font, c.number_format = _f(color=BLUE), "0.0"
        r += 1
    for st in STAGES:
        ws.cell(row=r, column=1, value="전결 처리").font = _f()
        ws.cell(row=r, column=2, value=st).font = _f()
        for j, g in enumerate(GRADES):
            c = ws.cell(row=r, column=3 + j, value=val.get(("전결 처리", st, g), 0.0))
            c.font, c.number_format = _f(color=BLUE), "0.0"
        stage_rows.append(r)
        r += 1
    last = r - 1
    total_row, sub_row, share_row = r, r + 1, r + 2
    labels = [(total_row, "합계 (주당)"), (sub_row, "전결 처리 소계"), (share_row, "전결 처리 비중")]
    for row, label in labels:
        ws.cell(row=row, column=1, value=label).font = _f(bold=True)
        for j in range(len(GRADES)):
            col = get_column_letter(3 + j)
            if row == total_row:
                f = f"=SUM({col}{first}:{col}{last})"
            elif row == sub_row:
                f = f"=SUM({col}{stage_rows[0]}:{col}{stage_rows[-1]})"
            else:
                f = f"=IF({col}{total_row}=0,0,{col}{sub_row}/{col}{total_row})"
            c = ws.cell(row=row, column=3 + j, value=f)
            c.font, c.fill = _f(bold=True), SUB_FILL
            c.number_format = "0.0%" if row == share_row else "0.0"
    for row in range(4, share_row + 1):
        for col in range(1, 3 + len(GRADES)):
            ws.cell(row=row, column=col).border = BOX
    # 비중 표: 각 칸 ÷ 직급군 합계
    base = share_row + 3
    ws.cell(row=base - 1, column=1, value="업무별 비중 (각 칸 ÷ 직급군 합계)").font = _f(bold=True, size=11)
    _head(ws, base, ["업무 범주", "세부 단계"] + GRADES)
    for i, src in enumerate(range(first, last + 1)):
        rr = base + 1 + i
        ws.cell(row=rr, column=1, value=f"=A{src}").font = _f()
        ws.cell(row=rr, column=2, value=f"=B{src}").font = _f()
        for j in range(len(GRADES)):
            col = get_column_letter(3 + j)
            c = ws.cell(row=rr, column=3 + j, value=f"=IF({col}${total_row}=0,0,{col}{src}/{col}${total_row})")
            c.font, c.number_format, c.border = _f(), "0.0%", BOX
        ws.cell(row=rr, column=1).border = ws.cell(row=rr, column=2).border = BOX
    ws.freeze_panes = "C5"
    _widths(ws, {"A": 16, "B": 18, "C": 11, "D": 11, "E": 11, "F": 11, "G": 12})


def _questions(ws) -> None:
    ws["A1"] = "가장 하기 싫은 업무 설문 (문항, 더미)"
    ws["A1"].font = _f(bold=True, size=13)
    _head(ws, 3, ["문항", "질문", "응답 형식", "선택지"])
    rows = [
        ("Q0", "소속 부서와 직급군", "선택", "부서: 연구1팀, 연구2팀, 시험평가팀, 경영지원팀 / 직급군: " + ", ".join(GRADES)),
        ("Q1", "결재·전결 업무 가운데 가장 하기 싫은 일 (1순위)", "단일 선택", ", ".join(OPTIONS)),
        ("Q2", "두 번째로 하기 싫은 일 (2순위)", "단일 선택 (Q1과 다른 것)", ", ".join(OPTIONS)),
        ("Q3", "세 번째로 하기 싫은 일 (3순위)", "단일 선택 (Q1·Q2와 다른 것)", ", ".join(OPTIONS)),
        ("Q4", "결재·전결 처리에 쓰는 주당 시간 (자가 응답)", "숫자 (시간)", "—"),
        ("Q5", "1순위 업무가 싫은 이유", "자유 응답", "—"),
    ]
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = _f(), BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=11, column=1, value="선택지 설명").font = _f(bold=True, size=11)
    _head(ws, 12, ["선택지", "뜻"])
    for i, st in enumerate(OPTIONS, 13):
        ws.cell(row=i, column=1, value=st).font = _f()
        ws.cell(row=i, column=2, value=STAGE_DESC.get(st, "결재와 무관한 기타 사무(자료 정리, 일정 조율 등)")).font = _f()
        ws.cell(row=i, column=1).border = ws.cell(row=i, column=2).border = BOX
    _widths(ws, {"A": 16, "B": 52, "C": 24, "D": 80})


def _responses(ws, survey: list[dict]) -> None:
    head = ["응답 ID", "부서", "직급군", "1순위 (Q1)", "2순위 (Q2)", "3순위 (Q3)", "주당 전결 처리 시간 (Q4)", "1순위 이유 (Q5)"]
    _head(ws, 1, head)
    keys = ["resp_id", "dept", "grade_group", "rank1", "rank2", "rank3", "hours_per_week_self", "reason"]
    for i, r in enumerate(survey, 2):
        for j, k in enumerate(keys, 1):
            v = float(r[k]) if k == "hours_per_week_self" else r[k]
            c = ws.cell(row=i, column=j, value=v)
            c.font = _f(color=BLUE)
            if k == "hours_per_week_self":
                c.number_format = "0.0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(survey) + 1}"
    _widths(ws, {"A": 9, "B": 12, "C": 10, "D": 15, "E": 15, "F": 15, "G": 14, "H": 40})


def _summary(ws, n: int) -> None:
    last = n + 1
    rng = lambda col: f"설문_응답!${col}$2:${col}${last}"  # noqa: E731
    ws["A1"] = "가장 하기 싫은 업무 설문 집계 (수식)"
    ws["A1"].font = _f(bold=True, size=13)
    ws["A2"] = "가중 점수 = 1순위 × 3 + 2순위 × 2 + 3순위 × 1. 응답을 고치면 다시 계산된다."
    ws["A2"].font = _f(color="52514E")
    ws["A3"] = "응답 수"
    ws["B3"] = f"=COUNTA({rng('A')})"
    ws["A3"].font = ws["B3"].font = _f(bold=True)
    _head(ws, 5, ["선택지", "1순위 수", "2순위 수", "3순위 수", "가중 점수", "순위", "1순위 비율"])
    top, bottom = 6, 6 + len(OPTIONS) - 1
    for i, opt in enumerate(OPTIONS, top):
        ws.cell(row=i, column=1, value=opt)
        ws.cell(row=i, column=2, value=f"=COUNTIF({rng('D')},$A{i})")
        ws.cell(row=i, column=3, value=f"=COUNTIF({rng('E')},$A{i})")
        ws.cell(row=i, column=4, value=f"=COUNTIF({rng('F')},$A{i})")
        ws.cell(row=i, column=5, value=f"=B{i}*3+C{i}*2+D{i}")
        ws.cell(row=i, column=6, value=f"=RANK(E{i},$E${top}:$E${bottom})")
        ws.cell(row=i, column=7, value=f"=IF($B$3=0,0,B{i}/$B$3)")
        for col in range(1, 8):
            c = ws.cell(row=i, column=col)
            c.font, c.border = _f(), BOX
        ws.cell(row=i, column=7).number_format = "0.0%"
    # 직급군별 1순위 교차표
    base = bottom + 3
    ws.cell(row=base - 1, column=1, value="직급군별 1순위 응답 수").font = _f(bold=True, size=11)
    _head(ws, base, ["선택지"] + GRADES)
    for i, opt in enumerate(OPTIONS, base + 1):
        ws.cell(row=i, column=1, value=opt).font = _f()
        ws.cell(row=i, column=1).border = BOX
        for j, g in enumerate(GRADES, 2):
            gl = get_column_letter(j)
            c = ws.cell(row=i, column=j, value=f"=COUNTIFS({rng('D')},$A{i},{rng('C')},{gl}${base})")
            c.font, c.border = _f(), BOX
    tr = base + 1 + len(OPTIONS)
    ws.cell(row=tr, column=1, value="응답 수").font = _f(bold=True)
    for j in range(2, 2 + len(GRADES)):
        gl = get_column_letter(j)
        c = ws.cell(row=tr, column=j, value=f"=SUM({gl}{base + 1}:{gl}{tr - 1})")
        c.font, c.fill, c.border = _f(bold=True), SUB_FILL, BOX
    ar = tr + 1
    ws.cell(row=ar, column=1, value="평균 주당 전결 처리 시간").font = _f(bold=True)
    for j in range(2, 2 + len(GRADES)):
        gl = get_column_letter(j)
        c = ws.cell(row=ar, column=j, value=f"=IFERROR(AVERAGEIF({rng('C')},{gl}${base},{rng('G')}),0)")
        c.font, c.fill, c.border, c.number_format = _f(bold=True), SUB_FILL, BOX, "0.0"
    _widths(ws, {"A": 22, "B": 11, "C": 11, "D": 11, "E": 11, "F": 9, "G": 11})
